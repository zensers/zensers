# -*- coding: utf-8 -*-
"""
Word 排版重构验收测试
=====================

构造最小 HTML → convert → 用 python-docx 重新打开输出文档，验证：
1. A4 页面与研报边距
2. 真 TOC 域与 updateFields 设置
3. 页脚 PAGE/NUMPAGES 域
4. H1 段前分页
5. 三线表（无 Table Grid、tblBorders、跨页重复表头 tblHeader）
6. 正文段首行缩进 w:firstLineChars="200"
7. 重复标题去重（H2 与前一 H1 去编号同文 → 跳过）
8. Normal 单倍行距 + 正文段显式固定 22pt 行距（行距下沉，防大字号裁切）
9. 标题自动编号（H1"一、" + 子标题相对深度编号 1.1/1.1.1，不依赖 h 标签固定层级）
"""

import os
import tempfile

import pytest
from PIL import Image
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

from src.converters.html_to_word import HTMLToWordConverter
from src.converters.word_styles import compute_image_width_cm, strip_heading_number


SAMPLE_HTML = """
<article>
    <div class="cover-page">
        <h1 class="cover-title">测试研究报告</h1>
        <p class="cover-badge">RESEARCH REPORT</p>
        <p class="cover-divider">━━━━━━━━━━</p>
        <p class="cover-meta">发布日期：2026年10月</p>
    </div>
    <div class="toc-page">
        <h2 class="toc-header">目 录</h2>
        <ul class="toc-list">
            <li class="toc-item toc-item-level1">01 产业链全景 3</li>
            <li class="toc-item toc-item-level2">01.1 核心结论 3</li>
        </ul>
    </div>
    <h1 class="chapter-title">01 产业链全景</h1>
    <h2 class="section-title">产业链全景</h2>
    <p>这是需要验证首行缩进的正文段落内容。</p>
    <table>
        <thead><tr><th>指标</th><th>数值</th></tr></thead>
        <tbody>
            <tr><td>市场规模</td><td>1.2万亿</td></tr>
            <tr><td>增长率</td><td>25%</td></tr>
        </tbody>
    </table>
</article>
"""


@pytest.fixture(scope="module")
def converted(tmp_path_factory):
    """转换一次并重新打开，供本模块所有断言复用"""
    out_dir = tmp_path_factory.mktemp("word_layout")
    output_path = os.path.join(str(out_dir), "layout.docx")
    converter = HTMLToWordConverter()
    result = converter.convert(SAMPLE_HTML, output_path)
    assert result.success is True
    return Document(output_path)


class TestWordPageSetup:
    """页面设置：A4 + 研报边距"""

    def test_a4_page_size(self, converted):
        section = converted.sections[0]
        assert abs(section.page_width.cm - 21.0) < 0.05
        assert abs(section.page_height.cm - 29.7) < 0.05

    def test_report_margins(self, converted):
        section = converted.sections[0]
        assert abs(section.top_margin.cm - 2.54) < 0.05
        assert abs(section.bottom_margin.cm - 2.54) < 0.05
        assert abs(section.left_margin.cm - 3.17) < 0.05
        assert abs(section.right_margin.cm - 3.17) < 0.05


class TestTocField:
    """目录：真 TOC 域 + updateFields"""

    def test_toc_field_exists(self, converted):
        body_xml = converted.element.xml
        assert 'TOC \\o' in body_xml

    def test_toc_placeholder_items_removed(self, converted):
        """目录区的手写目录条目不再输出为正文段落"""
        toc_texts = [p.text for p in converted.paragraphs if "核心结论" in p.text]
        assert toc_texts == []

    def test_update_fields_on_open(self, converted):
        settings_xml = converted.settings.element.xml
        assert "updateFields" in settings_xml
        assert 'w:val="true"' in settings_xml


class TestHeaderFooter:
    """页眉页脚：底边框 + PAGE/NUMPAGES 域"""

    def test_footer_contains_page_fields(self, converted):
        footer_xml = converted.sections[0].footer._element.xml
        assert "PAGE" in footer_xml
        assert "NUMPAGES" in footer_xml

    def test_footer_contains_page_number_text(self, converted):
        footer_texts = "".join(p.text for p in converted.sections[0].footer.paragraphs)
        assert "第 " in footer_texts and " 页 共 " in footer_texts

    def test_header_bottom_border(self, converted):
        header_para = converted.sections[0].header.paragraphs[0]
        p_pr = header_para._p.pPr
        assert p_pr is not None
        p_bdr = p_pr.find(qn("w:pBdr"))
        assert p_bdr is not None
        bottom = p_bdr.find(qn("w:bottom"))
        assert bottom is not None
        assert bottom.get(qn("w:sz")) == "4"
        assert bottom.get(qn("w:color")) == "1A2744"

    def test_header_contains_brand(self, converted):
        header_text = converted.sections[0].header.paragraphs[0].text
        assert "ZENSERS" in header_text


class TestHeadingTypography:
    """标题排版：H1 分页 + 重复标题去重"""

    def test_h1_page_break_before(self, converted):
        h1_paras = [p for p in converted.paragraphs if p.style.name == "Heading 1"]
        assert h1_paras, "应存在 Heading 1 段落"
        assert h1_paras[0].paragraph_format.page_break_before is True

    def test_h1_bottom_border(self, converted):
        h1_paras = [p for p in converted.paragraphs if p.style.name == "Heading 1"]
        p_bdr = h1_paras[0]._p.pPr.find(qn("w:pBdr"))
        assert p_bdr is not None
        bottom = p_bdr.find(qn("w:bottom"))
        assert bottom.get(qn("w:sz")) == "8"
        assert bottom.get(qn("w:color")) == "1A2744"

    def test_duplicate_h2_after_h1_skipped(self, converted):
        """<h1>01 产业链全景</h1><h2>产业链全景</h2> 只产出 1 个正文 heading（H1 重编号）"""
        headings = [p.text for p in converted.paragraphs if p.style.name.startswith("Heading")]
        assert headings == ["一、产业链全景"]

    def test_cover_title_not_styled_as_heading(self, converted):
        """封面主标题不占用 Heading 样式，也不更新去重状态"""
        cover_paras = [p for p in converted.paragraphs if p.text == "测试研究报告"]
        assert cover_paras and cover_paras[0].style.name != "Heading 1"


class TestBodyParagraph:
    """正文段落：首行缩进 2 字符"""

    def test_first_line_indent_chars(self, converted):
        body_paras = [p for p in converted.paragraphs if "首行缩进" in p.text]
        assert body_paras, "应存在正文段落"
        ind = body_paras[0]._p.pPr.find(qn("w:ind"))
        assert ind is not None
        assert ind.get(qn("w:firstLineChars")) == "200"
        assert ind.get(qn("w:firstLine")) == "420"

    def test_normal_style_simsun_font(self, converted):
        normal = converted.styles["Normal"]
        assert normal.font.size.pt == 10.5
        assert normal.font.name == "Times New Roman"
        east_asia = normal._element.rPr.rFonts.get(qn("w:eastAsia"))
        assert east_asia == "SimSun"


class TestThreeLineTable:
    """表格：三线表 + 跨页重复表头"""

    def test_no_table_grid_style(self, converted):
        table = converted.tables[0]
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblStyle")) is None
        assert "Table Grid" not in table._tbl.xml

    def test_tbl_borders_exist(self, converted):
        table = converted.tables[0]
        tbl_pr = table._tbl.tblPr
        borders = tbl_pr.find(qn("w:tblBorders"))
        assert borders is not None
        top = borders.find(qn("w:top"))
        assert top.get(qn("w:sz")) == "8"
        assert top.get(qn("w:color")) == "1A2744"
        assert borders.find(qn("w:insideV")).get(qn("w:val")) == "none"

    def test_header_row_repeat(self, converted):
        table = converted.tables[0]
        tr_pr = table.rows[0]._tr.trPr
        assert tr_pr is not None
        assert tr_pr.find(qn("w:tblHeader")) is not None

    def test_header_cell_shading(self, converted):
        table = converted.tables[0]
        tc_pr = table.rows[0].cells[0]._tc.tcPr
        shd = tc_pr.find(qn("w:shd"))
        assert shd is not None
        assert shd.get(qn("w:fill")) == "E8EDF4"


class TestWordStyleHelpers:
    """word_styles 工具函数单测"""

    def test_strip_heading_number_variants(self):
        assert strip_heading_number("01 产业链全景") == "产业链全景"
        assert strip_heading_number("第三章 市场规模") == "市场规模"
        assert strip_heading_number("一、供应商名录") == "供应商名录"
        assert strip_heading_number("01.2 子标题") == "子标题"
        assert strip_heading_number("无编号标题") == "无编号标题"

    def test_compute_image_width_wide(self):
        # 9×5.2in @200dpi = 1800×1040 → 宽度受 14.7cm 上限约束
        assert compute_image_width_cm(1800, 1040) == pytest.approx(14.7)

    def test_compute_image_width_small(self):
        # 800×400px @200dpi = 10.16cm（未触上限）
        assert compute_image_width_cm(800, 400) == pytest.approx(10.16, abs=0.01)

    def test_compute_image_width_square_capped(self):
        # 近方形（宽高比<1.3）上限 11cm
        assert compute_image_width_cm(1000, 900) == pytest.approx(11.0)


# ========== 用户验收缺陷回归 ==========
# 缺陷1：研究模板结构性标题泄漏进正文标题流
# 缺陷2：内联图片被 Normal 固定行距（EXACTLY 22pt）裁切成细条
# 缺陷3：三线表列宽均分难读、无斑马纹、行跨页断裂


def test_structural_headings_skipped():
    """结构性标题（论证与分析/核心结论等）不输出为正文 heading，内容保留"""
    html = ('<html><body><h1>01 测试章</h1>'
            '<h3>论证与分析</h3><p>正文内容A</p>'
            '<h3>核心结论</h3><p>正文内容B</p>'
            '<h3>合法小节</h3><p>正文内容C</p></body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    heads = [p.text.strip() for p in d.paragraphs if p.style.name.startswith('Heading')]
    assert not any('论证与分析' in h or '核心结论' in h for h in heads)
    assert any('合法小节' in h for h in heads)
    assert any('正文内容A' in p.text for p in d.paragraphs)  # 内容保留


def test_image_paragraph_line_spacing_reset():
    """图片段落行距重置为单倍，不继承 Normal 的 EXACTLY 22pt（避免裁切图片）"""
    tmp = tempfile.mkdtemp()
    img = os.path.join(tmp, 'ok.png')
    Image.new('RGB', (400, 300), 'white').save(img)
    html = f'<html><body><h1>T</h1><img src="{img}"/><p>tail</p></body></html>'
    out = os.path.join(tmp, 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    for p in d.paragraphs:
        if p._element.findall('.//' + qn('w:drawing')):
            pf = p.paragraph_format
            assert pf.line_spacing_rule in (None,) or str(pf.line_spacing_rule) != 'EXACTLY (4)'
            # 继承的固定行距在 paragraph_format 上读为 None，须断言段落级
            # 显式 spacing 且 lineRule="auto"（单倍行距），否则仍被裁切
            spacing = p._p.pPr.find(qn('w:spacing'))
            assert spacing is not None, '图片段落缺少显式行距，仍继承 Normal 固定 22pt'
            assert spacing.get(qn('w:lineRule')) == 'auto', spacing.get(qn('w:lineRule'))
            break
    else:
        pytest.fail('未找到图片段落')


def test_table_autofit_and_zebra():
    """三线表：autofit 布局 + 全宽，数据行隔行斑马纹 F7F9FC，行 cantSplit 防跨页断裂"""
    html = ('<html><body><h1>T</h1><table>'
            '<tr><th>指标</th><th>数值</th></tr>'
            '<tr><td>企业数量</td><td>140</td></tr>'
            '<tr><td>产品款数</td><td>330</td></tr>'
            '</table></body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    t = d.tables[0]
    assert t.autofit is True or t._tbl.tblPr.find(qn('w:tblLayout')) is not None

    def cell_fill(cell):
        shd = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
        return shd.get(qn('w:fill')) if shd is not None else None

    r1 = [cell_fill(c) for c in t.rows[1].cells]
    r2 = [cell_fill(c) for c in t.rows[2].cells]
    assert all(f == 'F7F9FC' for f in r1), r1
    assert all(f != 'F7F9FC' for f in r2), r2

    for row in t.rows:
        assert row._tr.trPr is not None and row._tr.trPr.find(qn('w:cantSplit')) is not None


# ========== 用户验收缺陷回归（第二轮） ==========
# 缺陷A：封面 28pt 标题/目录 18pt 标题被 Normal 固定行距（EXACTLY 22pt）裁切
#        → Normal 回归单倍，固定 22pt 只在正文段落显式设置（行距下沉）
# 缺陷B：正文标题无编号 → H1"一、" + 子标题相对深度编号（1.1/1.1.1，真 TOC 域自动带出）
# 缺陷C："图表 N"题注段落行距加固为显式单倍（保持无大纲级别，不进 TOC）


def test_body_exact_spacing_but_normal_single(converted):
    """Normal 恢复单倍行距；固定 22pt（exact 440 twips）只在正文段落显式设置"""
    # Normal 单倍（不再全局 EXACTLY，封面 28pt 标题/目录标题不再被裁切）
    assert converted.styles['Normal'].paragraph_format.line_spacing_rule != WD_LINE_SPACING.EXACTLY
    # 找一个正文段（非标题非列表非空）断言显式 exact 440 twips
    hit = False
    for p in converted.paragraphs:
        if p.style.name == 'Normal' and p.text.strip() and not p._element.findall('.//' + qn('w:drawing')):
            p_pr = p._element.find(qn('w:pPr'))
            sp = p_pr.find(qn('w:spacing')) if p_pr is not None else None
            if sp is not None and sp.get(qn('w:lineRule')) == 'exact':
                assert sp.get(qn('w:line')) == '440'
                hit = True
                break
    assert hit, '未找到显式 exact 22pt 行距的正文段落'


def test_heading_auto_numbering():
    """子标题相对深度编号：h4 第一层 → 1.1/1.2；h3 回浅级开新条目 1.3；
    其下 h4 → 1.3.1/1.3.2；再回 h3 → 1.4；新章 h4 → 2.1；
    前置区 H2 不编号；结构性标题跳过不消耗编号"""
    html = ('<html><body>'
            '<h2>前置说明</h2>'
            '<h1>行业概览</h1>'
            '<h4>定义与分类</h4>'
            '<h4>范围界定</h4>'
            '<h3>驱动因素</h3>'
            '<h4>政策驱动</h4>'
            '<h4>技术驱动</h4>'
            '<h3>抑制因素</h3>'
            '<h3>论证与分析</h3>'
            '<h1>竞争格局</h1>'
            '<h4>厂商梯队</h4>'
            '</body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    seq = [p.text.strip() for p in d.paragraphs if p.style.name.startswith('Heading')]
    assert seq == ['前置说明', '一、行业概览',
                   '1.1 定义与分类', '1.2 范围界定',      # h4 第一层
                   '1.3 驱动因素',                        # h3 回浅级 → 新的第一层条目
                   '1.3.1 政策驱动', '1.3.2 技术驱动',    # h4 深一层
                   '1.4 抑制因素',
                   '二、竞争格局', '2.1 厂商梯队'], seq


# ========== 用户验收缺陷回归（第六轮） ==========
# 缺陷：逻辑二级标题错用 Heading 4（编号已按相对深度计算，层级仍用 HTML 原始
#       标签 → h4 小节显示为 Heading 4，层级与编号脱节）
# 修复：Word 层级 = 相对深度 + 1（1.1 → Heading 2；1.1.1 → Heading 3；更深截断至 4）


def test_subheading_levels_normalized():
    """编号子标题的 Word 层级按相对深度归一：h4 首层小节 → Heading 2；
    其下 h5 嵌套子目 → Heading 3；正文 h2 回浅级 → Heading 2"""
    html = ('<html><body><h1>章</h1><h4>小节A</h4><h5>子目B</h5>'
            '<h2>关键发现</h2></body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    heads = {p.text.strip(): p.style.name for p in d.paragraphs
             if p.style.name.startswith('Heading')}
    assert heads.get('一、章') == 'Heading 1', heads
    assert heads.get('1.1 小节A') == 'Heading 2', heads
    assert heads.get('1.1.1 子目B') == 'Heading 3', heads
    kf = [(t, s) for t, s in heads.items() if '关键发现' in t]
    assert kf, heads
    assert kf[0][1] == 'Heading 2', kf


def test_prebody_heading_keeps_html_level():
    """前置区标题（h1c==0，不编号）保持 HTML 原层级"""
    html = ('<html><body><h3>摘要说明</h3><p>前置内容</p>'
            '<h1>正文章</h1><p>正文</p></body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    heads = {p.text.strip(): p.style.name for p in d.paragraphs
             if p.style.name.startswith('Heading')}
    assert heads.get('摘要说明') == 'Heading 3', heads


def test_chart_number_paragraph_single_spacing():
    """'图表 N' 题注段落显式单倍行距，且不设大纲级别（不进 TOC）"""
    tmp = tempfile.mkdtemp()
    img = os.path.join(tmp, 'ok.png')
    Image.new('RGB', (400, 300), 'white').save(img)
    html = f'<html><body><h1>T</h1><img src="{img}"/><p>tail</p></body></html>'
    out = os.path.join(tmp, 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    hit = False
    for p in d.paragraphs:
        if p.text.strip().startswith('图表 '):
            p_pr = p._p.pPr
            sp = p_pr.find(qn('w:spacing')) if p_pr is not None else None
            assert sp is not None, '图表题注段落缺少显式行距'
            assert sp.get(qn('w:lineRule')) == 'auto'
            assert p_pr.find(qn('w:outlineLvl')) is None
            hit = True
    assert hit, '未找到图表题注段落'


# ========== 用户验收缺陷回归（第三轮） ==========
# 缺陷1：div 嵌套覆盖 current_div_class → cover/toc 分页与样式分支全部失效（改 div 栈）
# 缺陷2：封面缺品牌元素（Logo + ZENSERS RESEARCH 字标 + 品牌尾分隔线/署名）
# 缺陷3：TOC 域漏掉 Heading 4 小节（1-3 → 1-4）


def test_cover_and_toc_page_breaks_with_nested_divs():
    html = ('<html><body>'
            '<div class="cover-page"><h1>测试报告标题</h1>'
            '<div class="meta-separator"></div><p class="meta">2026-10-11</p></div>'
            '<div class="toc"><h2>目 录</h2>'
            '<div class="toc-item"><a href="#s0">章一</a></div>'
            '<div class="toc-item"><a href="#s1">章二</a></div></div>'
            '<h1>一、正文章</h1><p>正文</p></body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    # 分页数：cover 尾 1 + toc 尾 1 = 至少 2
    brs = d.element.body.findall('.//' + qn('w:br'))
    page_brs = [b for b in brs if b.get(qn('w:type')) == 'page']
    assert len(page_brs) >= 2
    # toc-item 文本不重复输出（真 TOC 域已生成）
    texts = [p.text for p in d.paragraphs]
    assert not any(t.strip() == '章一' for t in texts)
    # 品牌元素
    assert any('ZENSERS RESEARCH' in p.text for p in d.paragraphs)
    assert any('行业深度研究报告' in p.text for p in d.paragraphs)
    # TOC 级别 1-4
    assert 'TOC \\o "1-4"' in d.element.xml


# ========== 用户验收缺陷回归（第四轮） ==========
# 缺陷：TOC 域更新后 Word 套用内置 TOC1-4 样式默认缩进（约 2 字符/层），
#       显式定义样式缩进，二级与一级缩进差控制在 1-2 字符


def test_toc_style_indents(converted):
    """TOC1-4 段落样式显式定义缩进（leftChars 字符单位），覆盖 Word 内置默认"""
    styles = {st.get(qn('w:styleId')): st
              for st in converted.styles.element.findall(qn('w:style'))}
    for sid, chars in (('TOC1', '0'), ('TOC2', '100'), ('TOC3', '200'), ('TOC4', '200')):
        st = styles.get(sid)
        assert st is not None, f'{sid} 未定义'
        p_pr = st.find(qn('w:pPr'))
        assert p_pr is not None, f'{sid} 缺少 pPr'
        ind = p_pr.find(qn('w:ind'))
        assert ind is not None and ind.get(qn('w:leftChars')) == chars, (sid, ind)


# ========== 内置封底（back cover 营销版式） ==========
# 缺陷：旧尾页（"关于 ZENSERS" + 4 条免责细则）内容单薄、字体大、全部堆在页首
# 要求：封底版式——内容沉底（顶部大留白）、全部居中、字号上限 13pt（主标语）、
#       免责压缩为底部一行法律小字；普通段落样式，不进 TOC


def test_ending_page_backcover_layout():
    """HTML 无 disclaimer 区块 → 追加封底（bottom-anchored 营销版式，不进 TOC）"""
    html = '<html><body><h1>T</h1><p>正文</p></body></html>'
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    paras = d.paragraphs

    def first_idx(pred):
        for i, p in enumerate(paras):
            if pred(p):
                return i
        return -1

    # 1) 自上而下顺序：字标 < 主标语 < 定位语 < 核心能力 < CTA < 法律行
    idx_mark = first_idx(lambda p: p.text.strip() == 'ZENSERS RESEARCH')
    idx_slogan = first_idx(lambda p: '自动化行业研究引擎' in p.text)
    idx_position = first_idx(lambda p: '让每一次行业洞察' in p.text)
    idx_caps = first_idx(lambda p: p.text.strip() == '核心能力')
    idx_cta = first_idx(lambda p: 'github.com/zensers/zensers' in p.text)
    idx_legal = first_idx(lambda p: '不构成投资建议' in p.text)
    assert 0 <= idx_mark < idx_slogan < idx_position < idx_caps < idx_cta < idx_legal, \
        (idx_mark, idx_slogan, idx_position, idx_caps, idx_cta, idx_legal)

    # 2) 顶部大留白空段（内容沉底）：space_before >= 380pt，位于字标之前
    spacers = [i for i, p in enumerate(paras[:idx_mark])
               if not p.text.strip()
               and p.paragraph_format.space_before is not None
               and p.paragraph_format.space_before.pt >= 380]
    assert spacers, '缺少封底顶部大留白空段（space_before >= 380pt）'

    # 3) 字号上限：封底所有 run 字号 <= 13pt
    for p in paras[spacers[0]:]:
        for r in p.runs:
            if r.font.size is not None:
                assert r.font.size.pt <= 13.0, (p.text, r.font.size.pt)

    # 4) 无 Heading 样式（不进 TOC）；可见段落全部居中；行距非 EXACTLY
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    for p in paras[spacers[0]:]:
        assert not p.style.name.startswith('Heading'), p.text
        if p.text.strip():
            assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER, p.text
        assert p.paragraph_format.line_spacing_rule != WD_LINE_SPACING.EXACTLY, p.text

    # 5) 封底特征文案：主标语、能力小标题、CTA、法律要件
    assert any('自动化行业研究引擎' in p.text for p in paras)
    assert any('核心能力' in p.text for p in paras)
    assert any('github.com/zensers/zensers' in p.text for p in paras)
    assert any('不构成投资建议' in p.text for p in paras)

    # 封底前有分页：文档 page break 计数至少 1
    brs = d.element.body.findall('.//' + qn('w:br'))
    page_brs = [b for b in brs if b.get(qn('w:type')) == 'page']
    assert len(page_brs) >= 1


def test_ending_page_not_duplicated_when_html_has_disclaimer():
    """HTML 已含 disclaimer-page 区块 → 不重复追加封底，自定义免责内容保留"""
    html = ('<html><body><h1>T</h1>'
            '<div class="disclaimer-page"><p>自定义免责内容</p></div>'
            '</body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    assert not any('自动化行业研究引擎' in t for t in texts)
    assert any('自定义免责内容' in t for t in texts)


# ========== 用户验收缺陷回归（第五轮） ==========
# 优化1：TOC1-4 样式行距略微扩大（约 1.4 倍 + 段前后 2pt）
# 优化2：封面布局重构——日期/署名沉底，副题与日期之间大留白空段


def test_toc_line_spacing_enlarged(converted):
    """TOC1 样式 pPr 含 spacing：line=340 lineRule=auto，段前后 40 twips（2pt）"""
    st = next(s for s in converted.styles.element.findall(qn('w:style'))
              if s.get(qn('w:styleId')) == 'TOC1')
    sp = st.find(qn('w:pPr') + '/' + qn('w:spacing'))
    assert sp is not None
    assert sp.get(qn('w:line')) == '340'
    assert sp.get(qn('w:lineRule')) == 'auto'
    assert sp.get(qn('w:before')) == '40'
    assert sp.get(qn('w:after')) == '40'


def test_cover_layout_order_and_bottom_meta():
    """封面自上而下：字标 < 标题 < 副题 < 日期 < 署名；副题与日期之间有
    space_before >= 200pt 的大留白空段；日期取自 meta 且只出现一次；
    封面段行距非 EXACTLY"""
    html = ('<html><body><div class="cover-page"><h1>测试报告标题</h1>'
            '<div class="meta-separator"></div><p class="meta">2026-10-11</p></div>'
            '<div class="toc"><h2>目 录</h2></div>'
            '<h1>一、正文章</h1><p>正文</p></body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    result = HTMLToWordConverter().convert(html, out)
    assert result.success
    d = Document(out)
    paras = d.paragraphs

    def first_idx(pred):
        for i, p in enumerate(paras):
            if pred(p):
                return i
        return -1

    idx_brand = first_idx(lambda p: p.text.strip() == 'ZENSERS RESEARCH')
    idx_title = first_idx(lambda p: p.text.strip() == '测试报告标题')
    idx_sub = first_idx(lambda p: '行业深度研究报告' in p.text)
    idx_date = first_idx(lambda p: p.text.strip() == '2026-10-11')
    idx_sign = first_idx(lambda p: 'zensers Research · github.com/zensers/zensers' in p.text)
    assert 0 <= idx_brand < idx_title < idx_sub < idx_date < idx_sign, \
        (idx_brand, idx_title, idx_sub, idx_date, idx_sign)

    # 大留白空段（副题与日期之间，space_before >= 200pt）
    spacers = [p for p in paras if not p.text.strip()
               and p.paragraph_format.space_before is not None
               and p.paragraph_format.space_before.pt >= 200]
    assert spacers, '缺少 space_before >= 200pt 的大留白空段'
    spacer_idx = paras.index(spacers[0])
    assert idx_sub < spacer_idx < idx_date, (idx_sub, spacer_idx, idx_date)

    # 日期文本来自 meta 且只出现一次（不再立即输出）
    date_hits = [p.text for p in paras if '2026-10-11' in p.text]
    assert len(date_hits) == 1, date_hits

    # 封面段行距非 EXACTLY（防大字号/留白段裁切）
    for p in paras[:idx_sign + 1]:
        rule = p.paragraph_format.line_spacing_rule
        assert rule != WD_LINE_SPACING.EXACTLY, (p.text, rule)


if __name__ == "__main__":
    pytest.main([__file__, "-v"])


def test_ordered_lists_restart_numbering_per_list():
    """每个 <ol> 从 1 重置编号，不再全文累加（修复第 N 个列表显示 51. 的缺陷）。"""
    html = ('<html><body><h1>一、章A</h1>'
            '<ol><li>甲项</li><li>乙项</li><li>丙项</li></ol>'
            '<h1>二、章B</h1>'
            '<ol><li>风险一</li><li>风险二</li></ol>'
            '</body></html>')
    out = os.path.join(tempfile.mkdtemp(), 't.docx')
    r = HTMLToWordConverter().convert(html, out)
    assert r.success
    d = Document(out)
    texts = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    assert any(x.startswith('1. 甲项') for x in texts)
    assert any(x.startswith('3. 丙项') for x in texts)
    assert any(x.startswith('1. 风险一') for x in texts), '第二个列表应重置为 1'
    assert not any(x.startswith('4. 风险一') for x in texts)
    # 不再依赖 List Number 样式（避免 Word 全文共享编号序列）
    assert all(p.style.name != 'List Number' for p in d.paragraphs)
