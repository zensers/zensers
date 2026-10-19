# -*- coding: utf-8 -*-
"""
Word 排版设计令牌与 OOxml 工具函数
==================================

为研究报告 Word 排版规范提供统一的设计令牌与底层 OOxml 操作，
供 html_to_word 转换器调用：

- 页面：A4（21 x 29.7cm），边距上下 2.54cm、左右 3.17cm
- 字体：标题黑体（SimHei）、正文宋体（SimSun）、西文 Times New Roman
- 色彩：藏青 NAVY / 金色 GOLD / 表头底纹等灰阶令牌

python-docx 1.2.0 无原生域（field）API，fldChar/instrText 等
OOxml 结构在此手工构造。
"""

import re

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ========== 色彩令牌 ==========
NAVY = "1A2744"                # 藏青：标题/三线表框线
GOLD = "C9A227"                # 金色：封面分隔线/强调
TABLE_HEADER_FILL = "E8EDF4"   # 表头浅藏青底纹
TABLE_ZEBRA_FILL = "F7F9FC"    # 数据行隔行斑马底纹
COVER_META_FILL = "F5F6F8"     # 封面元信息区浅灰底纹
GRAY_TEXT = "8C8C8C"           # 页眉页脚灰

# 研究写作模板的内部结构标题，不应进入正文标题流
STRUCTURAL_HEADINGS = {"论证与分析", "核心结论", "数据支撑", "风险提示",
                       "证据支撑", "证据链", "数据表格", "深度分析"}

# ========== 字体令牌 ==========
FONT_HEADING_CN = "SimHei"         # 标题黑体
FONT_BODY_CN = "SimSun"            # 正文宋体
FONT_WESTERN = "Times New Roman"   # 西文

# ========== 页面令牌 ==========
PAGE_WIDTH_CM = 21.0
PAGE_HEIGHT_CM = 29.7
MARGIN_TOP_BOTTOM_CM = 2.54
MARGIN_LEFT_RIGHT_CM = 3.17
CONTENT_WIDTH_CM = PAGE_WIDTH_CM - 2 * MARGIN_LEFT_RIGHT_CM  # ≈14.66cm

# ========== 图片规格 ==========
IMAGE_MAX_WIDTH_CM = 14.7       # 正文可用宽上限
IMAGE_SQUARE_MAX_CM = 11.0      # 饼/雷达等近方形图上限
IMAGE_ASSUME_DPI = 200          # 图表渲染 DPI 标准
IMAGE_FALLBACK_WIDTH_CM = 12.7  # PIL 读取失败时的回退宽度

# 剥离标题前导编号（"01 标题" / "第一章 标题" / "一、标题"）
_HEADING_NUMBER_PATTERN = re.compile(
    r"^\s*(第[一二三四五六七八九十百\d]+章|\d+(?:\.\d+)*|[一二三四五六七八九十]+、)\s*"
)


def strip_heading_number(text: str) -> str:
    """剥离标题前导编号，用于重复标题判断。

    "01 产业链全景" / "第一章 产业链全景" / "一、产业链全景" → "产业链全景"
    """
    return _HEADING_NUMBER_PATTERN.sub("", text or "", count=1).strip()


# ========== 标题自动编号 ==========
_CN_NUM = "一二三四五六七八九十"


def cn_num(n: int) -> str:
    """阿拉伯数 → 中文数字（H1 章序号用）：1→"一"、10→"十"、11→"十一"。

    仅处理 1..19 的中文数字习惯写法，≥20 回退阿拉伯数字；n 需 >= 1。
    """
    if n <= 10:
        return _CN_NUM[n - 1]
    if n < 20:
        return "十" + (_CN_NUM[n - 11] if n > 10 else "")
    return str(n)


def _make_border_element(tag: str, val: str, sz: int, color: str) -> "OxmlElement":
    """构造单个边框元素（w:top/w:bottom/w:left/...）。"""
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), color)
    return el


def style_three_line_table(table) -> None:
    """应用三线表边框：上下 1pt 藏青粗线，左右/内线 none。

    w:sz 单位为 1/8 pt：1pt → sz=8。
    表头行内单元格的 0.5pt 下边线由调用方配合 tcBorders 处理。
    """
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)

    borders = OxmlElement("w:tblBorders")
    borders.append(_make_border_element("top", "single", 8, NAVY))
    borders.append(_make_border_element("bottom", "single", 8, NAVY))
    for tag in ("left", "right", "insideH", "insideV"):
        borders.append(_make_border_element(tag, "none", 0, "auto"))

    # tblBorders 在 tblPr 子元素序列中位于 shd/tblLayout/tblCellMar/tblLook 之前
    tbl_pr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook"
    )


def mark_header_row_repeat(row) -> None:
    """给表头行加 w:trPr/w:tblHeader，跨页时自动重复表头。"""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def mark_row_cant_split(row) -> None:
    """给表格行加 w:trPr/w:cantSplit，防止行内容跨页断裂。

    trPr 子元素序列中 cantSplit 位于 trHeight/tblHeader 之前，
    需 insert_element_before 保持 schema 顺序（表头行已含 tblHeader）。
    """
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.insert_element_before(
            cant_split,
            "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc",
            "w:hidden", "w:ins", "w:del", "w:trPrChange",
        )


def add_field_run(paragraph, instr: str, placeholder: str = ""):
    """在段落内构造完整 Word 域结构，返回创建的 run 列表。

    结构：fldChar begin → instrText → fldChar separate → 占位 run → fldChar end。
    用于 TOC、PAGE、NUMPAGES 等域；占位 run 的文本在域刷新前显示。
    """
    runs = []

    r_begin = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    r_begin._r.append(fld_begin)
    runs.append(r_begin)

    r_instr = paragraph.add_run()
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    r_instr._r.append(instr_el)
    runs.append(r_instr)

    r_sep = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    r_sep._r.append(fld_sep)
    runs.append(r_sep)

    r_ph = paragraph.add_run(placeholder)
    runs.append(r_ph)

    r_end = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_end._r.append(fld_end)
    runs.append(r_end)

    return runs


def define_toc_styles(doc, indents=(0, 1, 2, 2)):
    """显式定义 TOC1-4 段落样式缩进（leftChars 字符单位），覆盖 Word 内置默认。

    TOC 域更新后 Word 会套用内置 TOC1-4 样式（默认约 2 字符/层缩进），
    此处预定义样式使二级与一级缩进差控制在 1-2 字符。

    indents[i] = TOC(i+1) 的左缩进字符数。
    """
    styles_el = doc.styles.element
    existing = {st.get(qn('w:styleId')) for st in styles_el.findall(qn('w:style'))}
    for i, chars in enumerate(indents, start=1):
        sid = f"TOC{i}"
        if sid in existing:
            st = next(s for s in styles_el.findall(qn('w:style'))
                      if s.get(qn('w:styleId')) == sid)
        else:
            st = OxmlElement('w:style')
            st.set(qn('w:type'), 'paragraph')
            st.set(qn('w:styleId'), sid)
            name = OxmlElement('w:name')
            name.set(qn('w:val'), f'toc {i}')
            st.append(name)
            based = OxmlElement('w:basedOn')
            based.set(qn('w:val'), 'Normal')
            st.append(based)
            styles_el.append(st)
        p_pr = st.find(qn('w:pPr'))
        if p_pr is None:
            p_pr = OxmlElement('w:pPr')
            st.append(p_pr)
        ind = p_pr.find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind')
            p_pr.append(ind)
        ind.set(qn('w:leftChars'), str(chars * 100))
        ind.set(qn('w:left'), str(chars * 210))  # twips 兜底（10.5pt ≈ 210/字符）

        # 行距略微扩大：约 1.4 倍行距（340/240）+ 段前后 2pt（40 twips）
        spacing = p_pr.find(qn('w:spacing'))
        if spacing is None:
            spacing = OxmlElement('w:spacing')
            # pPr 子元素序列：spacing 位于 ind 之前，须按 schema 顺序插入
            p_pr.insert_element_before(
                spacing,
                'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
                'w:suppressOverlap', 'w:jc', 'w:textDirection',
                'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
                'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange',
            )
        spacing.set(qn('w:line'), '340')
        spacing.set(qn('w:lineRule'), 'auto')
        spacing.set(qn('w:before'), '40')
        spacing.set(qn('w:after'), '40')


def set_update_fields_on_open(doc) -> None:
    """settings.xml 写入 w:updateFields val=true，Word/WPS 打开时自动刷新域。"""
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is not None:
        return
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    # updateFields 在 settings 子元素序列中位于 compat/rsids/mathPr 等之前
    settings.insert_element_before(
        el,
        "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
        "w:rsids", "m:mathPr", "w:themeFontLang", "w:clrSchemeMapping",
        "w:doNotAutoCompressPictures", "w:shapeDefaults",
        "w:decimalSymbol", "w:listSeparator",
    )


def set_paragraph_bottom_border(paragraph, sz: int, color: str) -> None:
    """给段落加底边框（w:pPr/w:pBdr/w:bottom）。

    Args:
        sz: 边框粗细，单位 1/8 pt（如 1pt → 8，0.5pt → 4，2.25pt → 18）
        color: 十六进制颜色（不带 #）
    """
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        # pBdr 在 pPr 序列中位于 shd/tabs/spacing/ind/jc 等之前
        p_pr.insert_element_before(
            p_bdr,
            "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
            "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
            "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
            "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
            "w:mirrorIndents", "w:suppressOverlap", "w:jc",
            "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
            "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
            "w:pPrChange",
        )
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)


def set_first_line_indent_chars(paragraph, chars: int = 200) -> None:
    """正文段落首行缩进（字符单位 + twips 兜底）。

    w:ind 同时写入 w:firstLineChars="200"（随字号缩放）与
    w:firstLine="420" twips（按 10.5pt 正文两字符兜底）。
    """
    from docx.shared import Pt

    # python-docx API 负责创建位置正确的 w:ind
    paragraph.paragraph_format.first_line_indent = Pt(21)  # 420 twips
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return
    ind = p_pr.find(qn("w:ind"))
    if ind is not None:
        ind.set(qn("w:firstLineChars"), str(chars))


def shade_cell_or_paragraph(obj, fill: str) -> None:
    """给表格单元格（_Cell）或段落（Paragraph）加底纹 w:shd。"""
    if hasattr(obj, "_tc"):
        pr = obj._tc.get_or_add_tcPr()
    elif hasattr(obj, "_p"):
        pr = obj._p.get_or_add_pPr()
    else:
        raise TypeError(f"unsupported object: {type(obj)!r}")

    shd = pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        if hasattr(obj, "_tc"):
            # tcPr 序列：shd 位于 noWrap/tcMar/textDirection/vAlign/hideMark 之前
            pr.insert_element_before(
                shd, "w:noWrap", "w:tcMar", "w:textDirection",
                "w:tcFitText", "w:vAlign", "w:hideMark",
            )
        else:
            # pPr 序列：shd 位于 tabs/spacing/ind/jc 等之前
            pr.insert_element_before(
                shd,
                "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                "w:suppressOverlap", "w:jc", "w:textDirection",
                "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
            )
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_table_header_cell_border(cell, sz: int = 4, color: str = NAVY) -> None:
    """表头单元格加 0.5pt 底边线（w:tcPr/w:tcBorders/w:bottom）。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        # tcBorders 在 tcPr 序列中位于 shd/noWrap/tcMar/vAlign 等之前
        tc_pr.insert_element_before(
            borders, "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
            "w:tcFitText", "w:vAlign", "w:hideMark",
        )
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), color)


def compute_image_width_cm(px_w: int, px_h: int) -> float:
    """按像素宽与 200dpi 计算图片插入宽度（cm），含上限规则。

    - 宽度 = min(14.7cm, 像素宽 / 200dpi)
    - 宽高比 < 1.3（饼/雷达等近方形图）上限 11cm
    """
    if px_w <= 0:
        return IMAGE_FALLBACK_WIDTH_CM
    width_cm = min(IMAGE_MAX_WIDTH_CM, px_w / float(IMAGE_ASSUME_DPI) * 2.54)
    if px_h > 0 and (px_w / float(px_h)) < 1.3:
        width_cm = min(width_cm, IMAGE_SQUARE_MAX_CM)
    return width_cm
