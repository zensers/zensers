# -*- coding: utf-8 -*-
"""
读 research_result.json（真实研究结果）→ Markdown转HTML → 锚点插入真实图表
→ 德勤模板 professional_report.html → HTMLToWordConverter → docx + 页数/证据验收。
图表 PNG 均来自持久目录（data/html_reports/charts），并汇总拷贝到案例目录。
"""
from dotenv import load_dotenv
load_dotenv()

import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

CASE_DIR = ROOT / "output" / "cases" / "humanoid_robot"
JSON_FILE = CASE_DIR / "research_result.json"
CHART_COPY_DIR = CASE_DIR / "charts"


def md_to_html(text: str) -> str:
    lines, out, in_table, rows = text.split("\n"), [], False, []

    def flush():
        nonlocal in_table, rows
        if len(rows) < 2:
            for r in rows:
                out.append("<p>" + "|".join(r) + "</p>")
        else:
            header = rows[0]
            n = len(header)
            body = []
            for r in rows[1:]:
                if all(set(c.strip()) <= set("- :") for c in r if c.strip()):
                    continue
                r = (r + [""] * n)[:n]
                body.append("<tr>" + "".join(f"<td>{re.sub(r'\*\*(.+?)\*\*', r'\1', c.strip())}</td>" for c in r) + "</tr>")
            out.append("<table><thead><tr>" + "".join(f"<th>{re.sub(r'\*\*(.+?)\*\*', r'\1', c.strip())}</th>" for c in header)
                       + "</tr></thead><tbody>" + "".join(body) + "</tbody></table>")
        in_table, rows = False, []

    for line in lines:
        s = line.strip()
        if not s:
            if in_table:
                flush()
            out.append("")
            continue
        if s.startswith("|") and s.endswith("|"):
            in_table = True
            rows.append([c.strip() for c in s.split("|")[1:-1]])
            continue
        if in_table:
            flush()
        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            t = s.lstrip("#").strip()
            if lvl == 1:
                out.append(f'<h2 class="section-title">{t}</h2>')
            elif lvl == 2:
                out.append(f'<h3 class="subsection-title">{t}</h3>')
            elif lvl == 3:
                out.append(f'<h4 class="sub-subsection-title">{t}</h4>')
            else:
                out.append(f"<p>{t}</p>")
            continue
        m_num = re.match(r"^(\d+)[.、]\s*(.+)$", s)
        m_bullet = re.match(r"^[-*]\s+(.+)$", s)
        if m_num:
            out.append("<li>" + m_num.group(2) + "</li>")
            continue
        if m_bullet:
            out.append("<li>" + m_bullet.group(1) + "</li>")
            continue
        s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
        s = re.sub(r"\*(.+?)\*", r"<em>\1</em>", s)
        out.append(f"<p>{s}</p>")
    if in_table:
        flush()
    return "\n".join(out)


def insert_charts(html: str, sections: list) -> tuple:
    """按 insertion_anchor/anchor_type 将真实图表插入章节 HTML，返回 (html, 嵌入数)"""
    count = 0
    for sec in sections:
        h1 = f'<h1 class="chapter-title">'
        end_h = f'</h1>'
        # 定位本章 h1 与下一章 h1 之间的片段
        marker = f'{sec.get("order", 0):02d} {sec.get("title", "")}'
        idx = html.find(marker)
        if idx < 0:
            continue
        start = html.rfind(h1, max(0, idx - 200), idx + len(marker))
        start = start if start >= 0 else max(0, idx - 60)
        nxt = html.find(h1, idx + len(marker))
        body = html[start:nxt] if nxt > 0 else html[start:]
        new_body = body
        for ch in sec.get("charts", []):
            src = ch.get("path", "")
            if not src or not Path(src).exists():
                print(f"  !! 图缺失跳过: {ch.get('title', '')}")
                continue
            cap_src = ch.get("subtitle", "")
            cap_note = ch.get("caption", "")
            fig = (f'<figure class="chart-block"><img src="{src}" alt="{ch.get("title", "")}"/>'
                   f'<figcaption><strong>{ch.get("title", "")}</strong>'
                   f'{("　数据来源：" + cap_src) if cap_src else ""}{("　" + cap_note) if cap_note else ""}</figcaption></figure>')
            anchor = ch.get("insertion_anchor", "")
            if anchor and anchor in new_body:
                pos = new_body.find("</p>", new_body.find(anchor))
                pos = pos + 4 if pos >= 0 else len(new_body)
            else:
                p0 = new_body.find("</p>")
                pos = p0 + 4 if p0 >= 0 else len(new_body)
            new_body = new_body[:pos] + fig + new_body[pos:]
            count += 1
        html = html[:start] + new_body + (html[nxt:] if nxt > 0 else "")
    return html, count


def main():
    report = json.loads(JSON_FILE.read_text(encoding="utf-8"))
    sections = report.get("sections", [])
    topic = report.get("topic", report.get("title", ""))
    template = (ROOT / "templates" / "professional_report.html").read_text(encoding="utf-8")

    toc, body = [], []
    for i, sec in enumerate(sections, 1):
        sec["order"] = i
        title = sec.get("title", "")
        toc.append(f'<li class="toc-item toc-item-level1"><span class="toc-num">{i:02d}</span>'
                   f'<span class="toc-title">{title}</span>'
                   f'<span class="toc-dots"></span>'
                   f'<span class="toc-page">{i + 2}</span></li>')
        # 二级：取 content 中 ## 标题，最多 4 条
        subs = []
        for line in str(sec.get("content", "")).splitlines():
            m = re.match(r"^##\s+(.+)$", line.strip())
            if m:
                t = re.sub(r"\*+", "", m.group(1)).strip()
                if t and t not in subs:
                    subs.append(t)
            if len(subs) >= 4:
                break
        for j, st in enumerate(subs, 1):
            toc.append(f'<li class="toc-item toc-item-level2"><span class="toc-num">{i:02d}.{j}</span>'
                       f'<span class="toc-title">{st}</span>'
                       f'<span class="toc-dots"></span>'
                       f'<span class="toc-page">{i + 2}</span></li>')
        body.append(f'<h1 class="chapter-title">{i:02d} {title}</h1>')
        body.append(f'<div class="section-content">{md_to_html(sec.get("content", ""))}</div>')

    findings = report.get("key_findings", [])
    kf = ""
    if findings:
        kf = '<div class="key-findings"><h2>关键发现</h2><ul>'
        for f in findings:
            kf += "<li>" + re.sub(r"\*\*", "", f) + "</li>"
        kf += "</ul></div>"

    html = template.replace("{{title}}", topic)
    html = html.replace("{{subtitle}}", "行业深度研究报告")
    html = html.replace("{{date}}", datetime.now().strftime("%Y年%m月"))
    html = html.replace("{{report_type}}", "行业深度研究")
    html = html.replace("{{toc_items}}", "\n".join(toc))
    html = html.replace("{{content}}", kf + "\n".join(body))
    html = html.replace("{{page_num}}", "1")

    html, embedded = insert_charts(html, sections)
    (CASE_DIR / "report.html").write_text(html, encoding="utf-8")

    CHART_COPY_DIR.mkdir(parents=True, exist_ok=True)
    copied = 0
    for sec in sections:
        for ch in sec.get("charts", []):
            p = Path(ch.get("path", ""))
            if p.exists():
                shutil.copy2(p, CHART_COPY_DIR / p.name)
                copied += 1

    docx_path = CASE_DIR / "中国人形机器人产业深度研究报告.docx"
    from src.converters.html_to_word import HTMLToWordConverter
    conv = HTMLToWordConverter().convert(html=html, output_path=str(docx_path))
    if not conv.success:
        print(f"转换失败: {conv.error}")
        sys.exit(1)

    from docx import Document
    import docx.oxml.ns as ns
    doc = Document(str(docx_path))
    chars = sum(len(p.text) for p in doc.paragraphs) + sum(len(c.text) for t in doc.tables for r in t.rows for c in r.cells)
    imgs = len(doc.element.body.findall(".//" + ns.qn("a:blip")))
    pages = max(1, chars // 2000) + imgs // 2

    src_cnt = len(report.get("sources", []))
    all_dps = [d for s in sections for d in s.get("data_points", [])]
    no_url = sum(1 for dp in all_dps if not dp.get("source_url"))
    da = report.get("defense_audit", {}) or {}
    l3_missing = sum(1 for i in (da.get("issues") or []) if i.get("code") == "missing_source_url")

    print("=" * 60)
    print(f"DOCX: {conv.output_path} ({conv.file_size / 1024:.0f} KB, 转换估页 {conv.pages_estimate})")
    print(f"自检: 字符 {chars} | 图表插入 {embedded}（docx嵌入 {imgs}，归档 {copied} PNG）| 表格 {len(doc.tables)}")
    print(f"自检: 估算页数 {pages} | 数据点 {len(all_dps)} 缺URL {no_url} | 源 {src_cnt} | 章节 {len(sections)}")
    print(f"验收: 页数>40 {'PASS' if pages > 40 else 'FAIL'} | 图>=12 {'PASS' if imgs >= 12 else 'FAIL'} "
          f"| L3缺URL=0 {'PASS' if l3_missing == 0 else f'FAIL({l3_missing})'} "
          f"| formal_complete={report.get('formal_complete')} | 质量门={report.get('quality_gate_status')}")
    print("=" * 60)


if __name__ == "__main__":
    main()